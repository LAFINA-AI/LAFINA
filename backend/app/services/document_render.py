"""Building the actual file from a checked document spec.

One function per format, each taking a spec that has already been through
document_guardrails.apply_output_limits and returning the file's bytes. They
are deterministic and never evaluate anything the model wrote: text is placed
as text, and the only markup any of them produce is bold.
"""

from __future__ import annotations

import io
import math
import re
from typing import Callable

from backend.app.services.document_guardrails import (
    MAX_SHEET_NAME_CHARS,
    is_allowed_formula,
    pdf_markup,
    plain_text,
    split_bold,
)
from backend.app.services.document_spec import DocumentSpec

AUTHOR = "LAFINA"

MIME_TYPES: dict[str, str] = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def _cell_text(value: object) -> str:
    """How a table value reads in a PDF or Word table."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


# ── PDF (ReportLab) ───────────────────────────────────────────────────────


def render_pdf(spec: DocumentSpec) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=plain_text(spec.title),
        author=AUTHOR,
        creator=AUTHOR,
    )
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    cell = ParagraphStyle("Cell", parent=body, fontSize=9, leading=11)
    header_cell = ParagraphStyle("HeaderCell", parent=cell, fontName="Helvetica-Bold")

    story: list = [Paragraph(pdf_markup(spec.title), styles["Title"])]
    for block in spec.blocks:
        if block.type == "heading":
            story.append(Paragraph(pdf_markup(block.text), styles[f"Heading{block.level}"]))
        elif block.type == "paragraph":
            story.append(Paragraph(pdf_markup(block.text), body))
        elif block.type in ("bullets", "numbered"):
            numbered = block.type == "numbered"
            story.append(
                ListFlowable(
                    [ListItem(Paragraph(pdf_markup(item), body)) for item in block.items],
                    bulletType="1" if numbered else "bullet",
                    start="1" if numbered else "•",
                    bulletFormat="%s." if numbered else None,
                    bulletFontName=body.fontName,
                    bulletFontSize=body.fontSize,
                    leftIndent=18,
                    spaceAfter=6,
                )
            )
        elif block.type == "table":
            width = len(block.columns) or max((len(row) for row in block.rows), default=0)
            if width == 0:
                continue
            data = []
            if block.columns:
                data.append([Paragraph(pdf_markup(column), header_cell) for column in block.columns])
            data.extend(
                [Paragraph(pdf_markup(_cell_text(value)), cell) for value in row]
                for row in block.rows
            )
            table = Table(
                data,
                colWidths=[document.width / width] * width,
                repeatRows=1 if block.columns else 0,
            )
            style = [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#C8C8C8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
            if block.columns:
                style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F0EEEA")))
            table.setStyle(TableStyle(style))
            story.extend([Spacer(1, 4), table, Spacer(1, 8)])

    document.build(story)
    return buffer.getvalue()


# ── Word (python-docx) ────────────────────────────────────────────────────


def _add_runs(paragraph, text: str) -> None:
    for run_text, bold in split_bold(text):
        run = paragraph.add_run(run_text)
        run.bold = bold or None


def render_docx(spec: DocumentSpec) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.core_properties.title = plain_text(spec.title)
    document.core_properties.author = AUTHOR
    document.add_heading(plain_text(spec.title), level=0)

    for block in spec.blocks:
        if block.type == "heading":
            document.add_heading(plain_text(block.text), level=block.level)
        elif block.type == "paragraph":
            _add_runs(document.add_paragraph(), block.text)
        elif block.type == "bullets":
            for item in block.items:
                _add_runs(document.add_paragraph(style="List Bullet"), item)
        elif block.type == "numbered":
            # Word's "List Number" style continues one count through the whole
            # document, so a second list would start at 4. Writing the number
            # keeps every list starting at 1.
            for index, item in enumerate(block.items, start=1):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Pt(18)
                paragraph.paragraph_format.first_line_indent = Pt(-18)
                paragraph.add_run(f"{index}.\t")
                _add_runs(paragraph, item)
        elif block.type == "table":
            width = len(block.columns) or max((len(row) for row in block.rows), default=0)
            if width == 0:
                continue
            table = document.add_table(rows=0, cols=width)
            try:
                table.style = "Light Grid Accent 1"
            except (KeyError, ValueError):
                table.style = "Table Grid"
            if block.columns:
                header = table.add_row().cells
                for index, column in enumerate(block.columns):
                    header[index].text = ""
                    run = header[index].paragraphs[0].add_run(plain_text(column))
                    run.bold = True
            for row in block.rows:
                cells = table.add_row().cells
                for index, value in enumerate(row):
                    cells[index].text = ""
                    _add_runs(cells[index].paragraphs[0], _cell_text(value))
            document.add_paragraph()

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── Excel (openpyxl) ──────────────────────────────────────────────────────

_SHEET_NAME_UNSAFE = re.compile(r"[\[\]:*?/\\]")
# No leading zeros: "0917 123 4567" and "007" are identifiers, not amounts.
_PLAIN_NUMBER = re.compile(r"^-?(0|[1-9]\d{0,14})(\.\d{1,10})?$")


def _sheet_name(name: str, taken: set[str], index: int) -> str:
    base = _SHEET_NAME_UNSAFE.sub(" ", plain_text(name or ""))
    base = " ".join(base.split()).strip("'")[:MAX_SHEET_NAME_CHARS] or f"Sheet{index}"
    candidate = base
    suffix = 2
    while candidate.lower() in taken:
        tail = f" ({suffix})"
        candidate = base[: MAX_SHEET_NAME_CHARS - len(tail)] + tail
        suffix += 1
    taken.add(candidate.lower())
    return candidate


def _write_cell(worksheet, row: int, column: int, value: object):
    """Writes one value, never letting model text become a formula it chose."""
    cell = worksheet.cell(row=row, column=column)
    if value is None:
        return cell
    if isinstance(value, bool):
        cell.value = value
    elif isinstance(value, (int, float)):
        cell.value = value if math.isfinite(value) else str(value)
    else:
        text = plain_text(str(value))
        if is_allowed_formula(text):
            cell.value = text.strip().upper()
        elif _PLAIN_NUMBER.match(text):
            # "42" and "3.50" are numbers the student will want to add up.
            number = float(text)
            cell.value = int(number) if number.is_integer() and "." not in text else number
        else:
            cell.value = text
            # openpyxl reads any string starting with "=" as a formula. Marking
            # the cell as text stores it literally, so =HYPERLINK(...) or a DDE
            # payload opens as the characters it is made of.
            cell.data_type = "s"
    return cell


def render_xlsx(spec: DocumentSpec) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = plain_text(spec.title)
    workbook.properties.creator = AUTHOR
    taken: set[str] = set()
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="F0EEEA")

    for index, sheet in enumerate(spec.sheets, start=1):
        worksheet = workbook.create_sheet(title=_sheet_name(sheet.name, taken, index))
        widths: dict[int, int] = {}
        row_number = 1
        if sheet.columns:
            for column_number, column in enumerate(sheet.columns, start=1):
                cell = _write_cell(worksheet, 1, column_number, column)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                widths[column_number] = len(str(cell.value or ""))
            worksheet.freeze_panes = "A2"
            row_number = 2
        for row in sheet.rows:
            for column_number, value in enumerate(row, start=1):
                cell = _write_cell(worksheet, row_number, column_number, value)
                if cell.value is not None and cell.data_type != "f":
                    widths[column_number] = max(
                        widths.get(column_number, 0), len(str(cell.value))
                    )
            row_number += 1
        for column_number, width in widths.items():
            worksheet.column_dimensions[get_column_letter(column_number)].width = min(
                60, max(10, width + 2)
            )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ── PowerPoint (python-pptx) ──────────────────────────────────────────────


# Every box is placed by hand on a 16:9 slide, in inches. The default
# template's layouts are 4:3, and a placeholder given any one coordinate of
# its own stops inheriting the rest — they fall to zero, which piles every box
# up at the top-left corner. So each box always gets all four.
SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 7.5

# The title slide: title and summary centred on the slide, a rule between them.
COVER_TITLE_BOX = (0.9, 1.6, 11.533, 2.1)
COVER_RULE_BOX = (6.067, 3.85, 1.2, 0.06)
COVER_SUMMARY_BOX = (1.667, 4.1, 10.0, 2.0)
# With no summary under it, the title and rule move down to the middle.
COVER_ALONE_SHIFT_IN = 0.6

# Content slides: a centred title band, a rule, and the bullets in a block
# below. The block is only as wide as its longest line and is centred on the
# slide, so a short list sits in the middle rather than against the left.
TITLE_BOX = (0.8, 0.4, 11.733, 1.15)
TITLE_RULE_BOX = (6.267, 1.62, 0.8, 0.05)
BODY_BOX = (1.667, 1.9, 10.0, 5.1)
BODY_MIN_WIDTH_IN = 4.0
BULLET_INDENT_IN = 0.4
# A slide with a title and nothing under it is a section divider.
DIVIDER_TITLE_BOX = (0.8, 2.6, 11.733, 2.3)

ACCENT_RGB = (0xF7, 0x5A, 0x5A)
SUMMARY_RGB = (0x59, 0x59, 0x59)

# Calibri, the template's font, averages under half an em per character in
# running text; this leaves room for bold, capitals and long words.
_CHAR_WIDTH_EM = 0.45
_LINE_HEIGHT = 1.2
# A text frame keeps 0.1" of padding on each side.
_FRAME_PADDING_IN = 0.2


def _estimated_height_in(
    texts: list[str], width_in: float, size_pt: int, space_after_pt: int = 0
) -> float:
    """How tall `texts` would stand at `size_pt` in a box `width_in` wide."""
    chars_per_line = max(1, int(width_in * 72 / (size_pt * _CHAR_WIDTH_EM)))
    lines = sum(max(1, math.ceil(len(text) / chars_per_line)) for text in texts)
    spacing = space_after_pt * max(0, len(texts) - 1)
    return (lines * size_pt * _LINE_HEIGHT + spacing) / 72


def _fitting_size(
    texts: list[str],
    box: tuple[float, float, float, float],
    sizes: tuple[int, ...],
    *,
    indent_in: float = 0.0,
    space_after_pt: int = 0,
) -> int:
    """The largest size in `sizes` at which `texts` fit the box; else the smallest."""
    _, _, width_in, height_in = box
    usable_width = width_in - _FRAME_PADDING_IN - indent_in
    # Padding, plus room to breathe: text that fills a box to its edge reads
    # as crammed even when it technically fits.
    usable_height = height_in - 0.4
    for size in sizes:
        if _estimated_height_in(texts, usable_width, size, space_after_pt) <= usable_height:
            return size
    return sizes[-1]


def _centred_to_content(
    texts: list[str],
    box: tuple[float, float, float, float],
    size_pt: int,
    *,
    indent_in: float = 0.0,
    min_width_in: float = 0.0,
) -> tuple[float, float, float, float]:
    """The box narrowed to its longest line at `size_pt`, centred on the slide.

    Generous by a sixth, so a line judged to fit never wraps for want of a
    few points.
    """
    _, top, width_in, height_in = box
    longest = max((len(text) for text in texts), default=0)
    needed = longest * size_pt * _CHAR_WIDTH_EM / 72 * 7 / 6 + indent_in + _FRAME_PADDING_IN
    width = min(width_in, max(min_width_in, needed))
    return ((SLIDE_WIDTH_IN - width) / 2, top, width, height_in)


def _shifted(box: tuple[float, float, float, float], down_in: float) -> tuple[float, float, float, float]:
    left, top, width, height = box
    return (left, top + down_in, width, height)


def _place(shape, box: tuple[float, float, float, float]) -> None:
    from pptx.util import Inches

    left, top, width, height = box
    shape.left, shape.top = Inches(left), Inches(top)
    shape.width, shape.height = Inches(width), Inches(height)


def _write_lines(
    shape,
    lines: list[str],
    size_pt: int,
    *,
    align,
    anchor,
    bold: bool = False,
    rgb: tuple[int, int, int] | None = None,
    space_after_pt: int = 0,
) -> None:
    """Replaces the shape's text with `lines`, one paragraph each, fully styled."""
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Pt

    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        if space_after_pt:
            paragraph.space_after = Pt(space_after_pt)
        for run_text, run_bold in split_bold(line):
            run = paragraph.add_run()
            run.text = run_text
            run.font.size = Pt(size_pt)
            if bold or run_bold:
                run.font.bold = True
            if rgb:
                run.font.color.rgb = RGBColor(*rgb)


def _add_rule(slide, box: tuple[float, float, float, float]) -> None:
    """A short accent bar, the one decoration the deck carries."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    left, top, width, height = box
    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = RGBColor(*ACCENT_RGB)
    rule.line.fill.background()
    rule.shadow.inherit = False


def _drop(shape) -> None:
    shape.element.getparent().remove(shape.element)


def render_pptx(spec: DocumentSpec) -> bytes:
    from pptx import Presentation
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)
    presentation.core_properties.title = plain_text(spec.title)
    presentation.core_properties.author = AUTHOR

    # ── Title slide ──
    cover = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_text = plain_text(spec.title)
    summary_text = plain_text(spec.summary)
    shift = 0.0 if summary_text else COVER_ALONE_SHIFT_IN
    cover_title_box = _shifted(COVER_TITLE_BOX, shift)
    _place(cover.shapes.title, cover_title_box)
    _write_lines(
        cover.shapes.title,
        [title_text],
        _fitting_size([title_text], cover_title_box, (44, 40, 36, 32, 28)),
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.BOTTOM,
        bold=True,
    )
    _add_rule(cover, _shifted(COVER_RULE_BOX, shift))
    subtitle = cover.placeholders[1] if len(cover.placeholders) > 1 else None
    if subtitle is not None:
        if summary_text:
            _place(subtitle, COVER_SUMMARY_BOX)
            _write_lines(
                subtitle,
                [summary_text],
                _fitting_size([summary_text], COVER_SUMMARY_BOX, (24, 22, 20, 18, 16, 14)),
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.TOP,
                rgb=SUMMARY_RGB,
            )
        else:
            _drop(subtitle)

    # ── Content slides ──
    for slide_spec in spec.slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        heading = plain_text(slide_spec.title)
        body = slide.placeholders[1]

        if not slide_spec.bullets:
            _drop(body)
            _place(slide.shapes.title, DIVIDER_TITLE_BOX)
            _write_lines(
                slide.shapes.title,
                [heading],
                _fitting_size([heading], DIVIDER_TITLE_BOX, (40, 36, 32, 28)),
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
                bold=True,
            )
        else:
            _place(slide.shapes.title, TITLE_BOX)
            _write_lines(
                slide.shapes.title,
                [heading],
                _fitting_size([heading], TITLE_BOX, (34, 30, 26, 22)),
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
                bold=True,
            )
            _add_rule(slide, TITLE_RULE_BOX)
            bullets = [plain_text(bullet) for bullet in slide_spec.bullets]
            space_after = 10
            size = _fitting_size(
                bullets,
                BODY_BOX,
                (28, 26, 24, 22, 20, 18, 16, 14, 12),
                indent_in=BULLET_INDENT_IN,
                space_after_pt=space_after,
            )
            _place(
                body,
                _centred_to_content(
                    bullets,
                    BODY_BOX,
                    size,
                    indent_in=BULLET_INDENT_IN,
                    min_width_in=BODY_MIN_WIDTH_IN,
                ),
            )
            _write_lines(
                body,
                slide_spec.bullets,
                size,
                align=PP_ALIGN.LEFT,
                anchor=MSO_ANCHOR.MIDDLE,
                space_after_pt=space_after,
            )

        if slide_spec.notes:
            slide.notes_slide.notes_text_frame.text = plain_text(slide_spec.notes)

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


RENDERERS: dict[str, Callable[[DocumentSpec], bytes]] = {
    "pdf": render_pdf,
    "docx": render_docx,
    "xlsx": render_xlsx,
    "pptx": render_pptx,
}
