"""Files the chat assistant writes: the spec, its guardrails, the renderers, the endpoint."""

import base64
import io
import json

import pytest
from httpx import AsyncClient

from backend.app.services.document_guardrails import (
    MAX_BLOCKS,
    MAX_SLIDES,
    apply_output_limits,
    is_allowed_formula,
    pdf_markup,
    safe_filename,
    split_bold,
)
from backend.app.services.document_render import RENDERERS
from backend.app.services.document_spec import (
    DocumentRefusal,
    DocumentSpecParseError,
    parse_document_spec,
)

# ── Specs a model might return ────────────────────────────────────────────

WRITTEN = {
    "title": "Cell Biology **Review**",
    "summary": "A one-page review of the cell organelles.",
    "blocks": [
        {"type": "heading", "text": "Organelles", "level": 1},
        {"type": "paragraph", "text": "The **mitochondrion** produces ATP."},
        {"type": "bullets", "items": ["Nucleus", "Ribosome"]},
        {"type": "numbered", "items": ["Read chapter 3", "Answer the quiz"]},
        {
            "type": "table",
            "columns": ["Organelle", "Job"],
            "rows": [["Nucleus", "Holds DNA"], {"Organelle": "Ribosome", "Job": "Makes protein"}],
        },
    ],
}

SHEETS = {
    "title": "Weekly Budget",
    "summary": "A budget with a running total.",
    "sheets": [
        {
            "name": "Budget: Sept/Oct",
            "columns": ["Item", "Cost"],
            "rows": [
                ["Food", 1200],
                ["Transport", "350"],
                ['=HYPERLINK("http://evil.example","click")', 0],
                ["Phone", "09171234567"],
                ["Total", "=SUM(B2:B5)"],
            ],
        }
    ],
}

SLIDES = {
    "title": "The Cell",
    "summary": "An introduction to cells.",
    "slides": [
        {"title": "What is a cell?", "bullets": ["The basic unit of life", "**Membrane** bound"]},
        {"title": "Organelles", "bullets": ["Nucleus", "Mitochondria"], "notes": "Pause here."},
    ],
}

SPECS = {"pdf": WRITTEN, "docx": WRITTEN, "xlsx": SHEETS, "pptx": SLIDES}


def _render(fmt: str, spec_dict: dict) -> tuple[bytes, list[str]]:
    spec = parse_document_spec(json.dumps(spec_dict), fmt)
    warnings = apply_output_limits(spec, fmt)
    return RENDERERS[fmt](spec), warnings


# ── Reading the model's reply ─────────────────────────────────────────────


def test_a_fenced_reply_with_extra_keys_still_reads():
    raw = "```json\n" + json.dumps({**WRITTEN, "author": "someone"}) + "\n```"
    spec = parse_document_spec(raw, "pdf")
    assert spec.title == "Cell Biology **Review**"
    table = spec.blocks[-1]
    # A row written as an object is read in column order.
    assert table.rows[1] == ["Ribosome", "Makes protein"]


def test_a_reply_without_content_for_the_format_is_refused():
    with pytest.raises(DocumentSpecParseError):
        parse_document_spec(json.dumps(WRITTEN), "pptx")
    with pytest.raises(DocumentSpecParseError):
        parse_document_spec("not json at all", "pdf")
    with pytest.raises(DocumentSpecParseError):
        parse_document_spec(json.dumps({"title": "x", "blocks": [{"type": "image"}]}), "pdf")


def test_a_refusal_carries_its_reason():
    with pytest.raises(DocumentRefusal) as refusal:
        parse_document_spec('{"refusal": "That would be a forged certificate."}', "pdf")
    assert refusal.value.reason == "That would be a forged certificate."


# ── Rendering ─────────────────────────────────────────────────────────────


def test_a_pdf_opens_and_reads_back():
    import pdfplumber

    data, _ = _render("pdf", WRITTEN)
    assert data.startswith(b"%PDF-")
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        text = pdf.pages[0].extract_text()
    assert "Cell Biology Review" in text
    assert "Holds DNA" in text


def test_a_word_file_opens_with_its_headings_and_table():
    import docx

    data, _ = _render("docx", WRITTEN)
    document = docx.Document(io.BytesIO(data))
    headings = [p.text for p in document.paragraphs if p.style.name.startswith(("Title", "Heading"))]
    assert headings == ["Cell Biology Review", "Organelles"]
    assert document.tables[0].rows[2].cells[0].text == "Ribosome"
    # Every numbered list starts at one.
    assert "1.\tRead chapter 3" in [p.text for p in document.paragraphs]


def test_a_workbook_opens_with_its_sheet_and_header():
    import openpyxl

    data, _ = _render("xlsx", SHEETS)
    workbook = openpyxl.load_workbook(io.BytesIO(data))
    assert workbook.sheetnames == ["Budget Sept Oct"]
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == ["Item", "Cost"]
    assert sheet.freeze_panes == "A2"
    # A number written as text is still a number, but an identifier keeps its zero.
    assert sheet["B3"].value == 350
    assert sheet["B5"].value == "09171234567"


def test_a_deck_opens_with_a_title_slide_and_notes():
    from pptx import Presentation

    data, _ = _render("pptx", SLIDES)
    deck = Presentation(io.BytesIO(data))
    assert [slide.shapes.title.text for slide in deck.slides] == [
        "The Cell",
        "What is a cell?",
        "Organelles",
    ]
    assert deck.slides[2].notes_slide.notes_text_frame.text == "Pause here."


# ── Guardrails ────────────────────────────────────────────────────────────


def test_only_simple_totals_survive_as_formulas():
    import openpyxl

    assert is_allowed_formula("=SUM(B2:B5)")
    assert is_allowed_formula("=average(c2:c9)")
    assert not is_allowed_formula('=HYPERLINK("http://x","click")')
    assert not is_allowed_formula("=WEBSERVICE(A1)")
    assert not is_allowed_formula("=SUM(B2:B5)+cmd|' /C calc'!A0")

    data, _ = _render("xlsx", SHEETS)
    sheet = openpyxl.load_workbook(io.BytesIO(data)).active
    assert sheet["A4"].data_type == "s"
    assert sheet["A4"].value.startswith("=HYPERLINK")
    assert sheet["B6"].data_type == "f"
    assert sheet["B6"].value == "=SUM(B2:B5)"


def test_pdf_markup_from_the_model_is_printed_not_obeyed():
    import pdfplumber

    assert pdf_markup('<img src="/etc/passwd"/> **bold**') == (
        '&lt;img src="/etc/passwd"/&gt; <b>bold</b>'
    )
    spec = {
        "title": "Notes",
        "blocks": [{"type": "paragraph", "text": '<img src="/etc/passwd" width="10"/> done'}],
    }
    data, _ = _render("pdf", spec)
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        assert '<img src="/etc/passwd" width="10"/> done' in pdf.pages[0].extract_text()


def test_characters_xml_cannot_hold_are_dropped_before_writing():
    import docx

    spec = {"title": "Bad\x00 bytes\x0b", "blocks": [{"type": "paragraph", "text": "a\x01b"}]}
    data, _ = _render("docx", spec)
    document = docx.Document(io.BytesIO(data))
    assert document.paragraphs[0].text == "Bad bytes"
    assert document.paragraphs[1].text == "ab"


def test_an_oversized_spec_is_cut_short_with_a_warning():
    long_doc = {
        "title": "Long",
        "blocks": [{"type": "paragraph", "text": f"Paragraph {n}"} for n in range(MAX_BLOCKS + 30)],
    }
    spec = parse_document_spec(json.dumps(long_doc), "pdf")
    warnings = apply_output_limits(spec, "pdf")
    assert len(spec.blocks) == MAX_BLOCKS
    assert warnings and str(MAX_BLOCKS) in warnings[0]

    long_deck = {"title": "Deck", "slides": [{"title": f"S{n}"} for n in range(MAX_SLIDES + 5)]}
    spec = parse_document_spec(json.dumps(long_deck), "pptx")
    assert apply_output_limits(spec, "pptx")
    assert len(spec.slides) == MAX_SLIDES


def test_filenames_are_safe_everywhere():
    assert safe_filename("Cell **Biology**: Week 1/2?", "pdf") == "Cell Biology Week 1 2.pdf"
    assert safe_filename("", "xlsx") == "LAFINA document.xlsx"
    assert safe_filename("CON", "docx") == "LAFINA document.docx"
    assert safe_filename("x" * 200, "pptx") == "x" * 60 + ".pptx"


def test_bold_markers_split_into_runs():
    assert split_bold("a **b** c") == [("a ", False), ("b", True), (" c", False)]
    assert split_bold("unclosed **bold") == [("unclosed bold", False)]


# ── The endpoint ──────────────────────────────────────────────────────────


class FakeDeepSeek:
    """Stands in for the provider: records calls, replies how it is told to."""

    def __init__(self, replies=None, error=None):
        self.replies = replies or [json.dumps(WRITTEN)]
        self.error = error
        self.calls = []

    async def json_completion(
        self, messages, user_id, request_id="", model=None, max_tokens=4096, json_mode=False
    ):
        self.calls.append(
            {"messages": messages, "model": model, "json_mode": json_mode, "request_id": request_id}
        )
        if self.error:
            raise self.error
        reply = self.replies[(len(self.calls) - 1) % len(self.replies)]
        return reply, {"prompt_tokens": 10, "completion_tokens": 5, "total_tokens": 15}


async def _headers(
    async_client: AsyncClient,
    email: str,
    *,
    role: str = "student_pro",
    plan: str = "student",
    system_role: str = "user",
) -> dict[str, str]:
    from sqlalchemy import select

    from backend.app.models.account import Account
    from backend.tests.conftest import TestingSessionLocal

    password = "super-strong-lafina-passphrase-2026"
    await async_client.post("/v1/auth/register", json={"email": email, "password": password})
    async with TestingSessionLocal() as db:
        account = (await db.execute(select(Account).where(Account.email == email))).scalar_one()
        account.role = role
        account.subscription_plan = plan
        account.system_role = system_role
        await db.commit()
    login = await async_client.post("/v1/auth/login", json={"email": email, "password": password})
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


def _body(fmt: str = "pdf", text: str = "Make me a PDF review of cell organelles.") -> dict:
    return {
        "format": fmt,
        "messages": [
            {"role": "user", "content": "What does the mitochondrion do?"},
            {"role": "assistant", "content": "It produces ATP."},
            {"role": "user", "content": text},
        ],
    }


@pytest.fixture
def fake_deepseek():
    from backend.app.api.v1.ai import get_deepseek_client
    from backend.app.main import app

    fakes: list[FakeDeepSeek] = []

    def install(fake: FakeDeepSeek) -> FakeDeepSeek:
        app.dependency_overrides[get_deepseek_client] = lambda: fake
        fakes.append(fake)
        return fake

    yield install
    app.dependency_overrides.pop(get_deepseek_client, None)


@pytest.mark.asyncio
@pytest.mark.parametrize("fmt", ["pdf", "docx", "xlsx", "pptx"])
async def test_a_student_pro_gets_a_file(async_client: AsyncClient, fake_deepseek, fmt):
    fake = fake_deepseek(FakeDeepSeek(replies=[json.dumps(SPECS[fmt])]))
    headers = await _headers(async_client, f"doc_{fmt}@ustp.edu.ph")

    res = await async_client.post("/v1/ai/documents", headers=headers, json=_body(fmt))
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["format"] == fmt
    assert data["filename"].endswith(f".{fmt}")
    content = base64.b64decode(data["contentBase64"])
    assert data["sizeBytes"] == len(content)
    assert content[:4] == (b"%PDF" if fmt == "pdf" else b"PK\x03\x04")
    assert data["summary"]
    assert data["usage"]["total_tokens"] == 15

    call = fake.calls[0]
    assert call["json_mode"] is True
    assert call["model"] == "deepseek-v4-flash"
    # The conversation reaches the model after the system instruction.
    assert call["messages"][0]["role"] == "system"
    assert call["messages"][-1]["content"].startswith("Make me a PDF")
    assert len(call["messages"]) == 4


@pytest.mark.asyncio
async def test_student_pro_by_plan_alone_is_enough(async_client: AsyncClient, fake_deepseek):
    fake_deepseek(FakeDeepSeek())
    headers = await _headers(async_client, "doc_plan@ustp.edu.ph", role="student", plan="student_pro")
    res = await async_client.post("/v1/ai/documents", headers=headers, json=_body())
    assert res.status_code == 200, res.text


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("role", "plan", "system_role"),
    [
        ("student", "student", "user"),
        ("business", "business", "user"),
        ("admin", "student", "admin"),
    ],
)
async def test_files_are_exclusive_to_student_pro(
    async_client: AsyncClient, fake_deepseek, role, plan, system_role
):
    fake = fake_deepseek(FakeDeepSeek())
    headers = await _headers(
        async_client, f"doc_{role}@ustp.edu.ph", role=role, plan=plan, system_role=system_role
    )
    res = await async_client.post("/v1/ai/documents", headers=headers, json=_body())
    assert res.status_code == 403
    assert "student_pro" in res.json()["detail"]
    assert fake.calls == []


@pytest.mark.asyncio
async def test_signing_in_is_required(async_client: AsyncClient):
    res = await async_client.post("/v1/ai/documents", json=_body())
    assert res.status_code == 401


@pytest.mark.asyncio
async def test_a_disabled_student_pro_is_refused(async_client: AsyncClient, fake_deepseek):
    from sqlalchemy import select

    from backend.app.models.account import Account
    from backend.tests.conftest import TestingSessionLocal

    fake = fake_deepseek(FakeDeepSeek())
    headers = await _headers(async_client, "doc_disabled@ustp.edu.ph")
    async with TestingSessionLocal() as db:
        account = (
            await db.execute(select(Account).where(Account.email == "doc_disabled@ustp.edu.ph"))
        ).scalar_one()
        account.is_active = False
        await db.commit()

    res = await async_client.post("/v1/ai/documents", headers=headers, json=_body())
    assert res.status_code in (401, 403)
    assert fake.calls == []


@pytest.mark.asyncio
async def test_a_request_is_checked_before_the_model_is_called(
    async_client: AsyncClient, fake_deepseek
):
    fake = fake_deepseek(FakeDeepSeek())
    headers = await _headers(async_client, "doc_checks@ustp.edu.ph")

    blank = await async_client.post("/v1/ai/documents", headers=headers, json=_body(text="   "))
    assert blank.status_code == 400

    wrong = await async_client.post("/v1/ai/documents", headers=headers, json=_body(fmt="exe"))
    assert wrong.status_code == 422

    long = _body()
    long["messages"] = [{"role": "user", "content": "x" * 4000}] * 3
    too_long = await async_client.post("/v1/ai/documents", headers=headers, json=long)
    assert too_long.status_code == 400

    assert fake.calls == []


@pytest.mark.asyncio
async def test_the_file_allowance_is_its_own(async_client: AsyncClient, fake_deepseek):
    fake_deepseek(FakeDeepSeek())
    headers = await _headers(async_client, "doc_quota@ustp.edu.ph")
    for _ in range(3):
        res = await async_client.post("/v1/ai/documents", headers=headers, json=_body())
        assert res.status_code == 200, res.text
    res = await async_client.post("/v1/ai/documents", headers=headers, json=_body())
    assert res.status_code == 429


@pytest.mark.asyncio
async def test_an_unreadable_reply_is_retried_once(async_client: AsyncClient, fake_deepseek):
    fake = fake_deepseek(FakeDeepSeek(replies=["Sure! Here is your PDF.", json.dumps(WRITTEN)]))
    headers = await _headers(async_client, "doc_retry@ustp.edu.ph")

    res = await async_client.post("/v1/ai/documents", headers=headers, json=_body())
    assert res.status_code == 200, res.text
    assert len(fake.calls) == 2
    # The second attempt saw its own bad reply and the nudge after it.
    retry = fake.calls[1]["messages"]
    assert retry[-2] == {"role": "assistant", "content": "Sure! Here is your PDF."}
    assert "JSON object" in retry[-1]["content"]
    assert res.json()["usage"]["total_tokens"] == 30


@pytest.mark.asyncio
async def test_two_unreadable_replies_are_a_bad_gateway(async_client: AsyncClient, fake_deepseek):
    fake = fake_deepseek(FakeDeepSeek(replies=["no json here"]))
    headers = await _headers(async_client, "doc_502@ustp.edu.ph")

    res = await async_client.post("/v1/ai/documents", headers=headers, json=_body())
    assert res.status_code == 502
    assert len(fake.calls) == 2


@pytest.mark.asyncio
async def test_a_refusal_is_passed_back(async_client: AsyncClient, fake_deepseek):
    fake_deepseek(FakeDeepSeek(replies=['{"refusal": "That would be a forged diploma."}']))
    headers = await _headers(async_client, "doc_refuse@ustp.edu.ph")

    res = await async_client.post("/v1/ai/documents", headers=headers, json=_body())
    assert res.status_code == 422
    assert "forged diploma" in res.json()["detail"]


@pytest.mark.asyncio
async def test_a_provider_failure_keeps_its_status(async_client: AsyncClient, fake_deepseek):
    from backend.app.clients.deepseek import DeepSeekBillingError

    fake_deepseek(FakeDeepSeek(error=DeepSeekBillingError()))
    headers = await _headers(async_client, "doc_503@ustp.edu.ph")

    res = await async_client.post("/v1/ai/documents", headers=headers, json=_body())
    assert res.status_code == 503
